from __future__ import annotations
import json
from io import BytesIO
from typing import Any
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import annotations as an
from models import AudioChunk, AudioInfo, TranscriptSegment
import strings as s
from speaker_reconciliation import apply_speaker_mapping
from transcription import MODEL

APP_VERSION = "1.1.2"


def format_timestamp(seconds: float) -> str:
    total=max(0,int(seconds)); hours, rem=divmod(total,3600); minutes, secs=divmod(rem,60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def readable_transcript(segments: list[TranscriptSegment], mapping: dict[str,str], timestamps: bool=True) -> str:
    blocks=[]
    for item in sorted(segments,key=lambda s:(s.absolute_start,s.absolute_end,s.chunk_index)):
        stamp=f" [{format_timestamp(item.absolute_start)}–{format_timestamp(item.absolute_end)}]" if timestamps else ""
        blocks.append(f"{apply_speaker_mapping(item,mapping)}{stamp}\n{item.text}")
    return "\n\n".join(blocks)


def project_payload(info: AudioInfo | None, chunks: list[AudioChunk], segments: list[TranscriptSegment], mapping: dict[str,str],
                    generated_at: str, metadata: dict[str,Any] | None=None, include_source_path: bool=False,
                    workflow_step: int=0, provider: str="", model: str="",
                    range_start: float|None=None, range_end: float|None=None,
                    last_reviewed: int|None=None, highlight_labels: list[str] | None=None,
                    reviewer: str="") -> dict[str,Any]:
    # The real provenance of the run; without it, any non-OpenAI path would be recorded
    # wrongly. The range transcribed is part of that provenance: without it there is no
    # way to know which portion of the recording was actually analysed.
    audio = info.to_dict() if info else None
    if audio and not include_source_path: audio["path"] = ""
    return {"application_version":APP_VERSION,"project_format":"transcript-project-v1","source_filename":info.filename if info else "",
            "source_path":info.path if info and include_source_path else None,"source_duration":info.duration if info else 0,
            "audio_metadata":audio,
            # A reference to the recording, so a reopened project can rebind its player.
            # The absolute path follows the same privacy rule as source_path; the filename and
            # size always travel, because they are what lets a moved file be re-identified.
            "audio_path":info.path if info and include_source_path else "",
            "audio_filename":info.filename if info else "",
            "audio_bytes":info.size_bytes if info else 0,
            "transcription_provider":provider or "openai","transcription_model":model or MODEL,
            "transcription_range_start":range_start,"transcription_range_end":range_end,"generation_date":generated_at,
            "chunks":[c.to_dict() for c in chunks],"speaker_mapping":mapping,
            "segments":[s.to_dict(apply_speaker_mapping(s,mapping)) for s in segments],
            # Review progress. Per-turn `checked` travels inside each segment; this records
            # only where the researcher was, so reopening puts them back at the same turn.
            "review":{"last_reviewed_turn_index":last_reviewed,
                      "checked_turns":sum(1 for item in segments if item.checked),
                      # Who was reviewing, so a reopened project keeps signing its comments
                      # with the same name rather than with whoever opens it next.
                      "reviewer":reviewer or ""},
            # What the highlight colours mean in this study. Saved with the project, because
            # a colour code that is lost when the file is reopened is not a code.
            "highlight_legend":list(highlight_labels or []),
            "interview_metadata":metadata or {},"workflow_step":workflow_step}


def make_json(*args: Any, **kwargs: Any) -> bytes:
    return json.dumps(project_payload(*args,**kwargs),ensure_ascii=False,indent=2).encode("utf-8")


def _paint(run: Any, slot: str) -> None:
    """Put a slot's colour behind a run, with the ink that stays readable on it.

    Word's highlighter has fifteen fixed colours and none of the three is among them, so
    every slot is written as run shading — a plain `w:shd` fill, which takes any RGB and
    renders as the same coloured band behind the words. The ink is set alongside it because
    Smoky Rose is dark enough that Word's default black on it is 3.1:1.
    """
    fill,ink=an.HIGHLIGHTS[slot]
    shading=OxmlElement("w:shd")
    shading.set(qn("w:val"),"clear"); shading.set(qn("w:color"),"auto"); shading.set(qn("w:fill"),fill)
    run._r.get_or_add_rPr().append(shading)
    run.font.color.rgb=RGBColor.from_string(ink)


def _legend(document: Any, segments: list[TranscriptSegment], labels: list[str] | None) -> None:
    """Say what the highlight colours mean, but only for the ones actually used.

    A colour code is only a code if the reader is told the key. Printing all three when the
    transcript uses one would be noise, so the legend lists what is on the page.
    """
    used={a.slot for item in segments for a in getattr(item,"annotations",None) or []
          if a.kind==an.HIGHLIGHT and a.slot in an.HIGHLIGHTS}
    if not used: return
    names=list(labels or [])
    line=document.add_paragraph(); line.add_run(f"{s.DOC_LEGEND}: ").bold=True
    for position,slot in enumerate(sorted(used)):
        if position: line.add_run("   ")
        _paint(line.add_run("  "),slot)
        index=int(slot)-1
        label=names[index] if index < len(names) and (names[index] or "").strip() else               s.DOC_LEGEND_UNNAMED.format(number=slot)
        line.add_run(f" {label}").font.size=Pt(10)


def _page_number(paragraph: Any) -> None:
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=paragraph.add_run()
    begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    instruction=OxmlElement("w:instrText"); instruction.set(qn("xml:space"),"preserve"); instruction.text=" PAGE "
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end"); run._r.extend([begin,instruction,end])


def make_docx(info: AudioInfo, segments: list[TranscriptSegment], mapping: dict[str,str], generated_at: str,
              title: str="", metadata: dict[str,str] | None=None,
              include_timestamps: bool=True, include_notice: bool=True, model: str="",
              comment_author: str="", highlight_labels: list[str] | None=None) -> bytes:
    try:
        document=Document(); section=document.sections[0]
        section.top_margin=section.bottom_margin=Cm(2.5); section.left_margin=section.right_margin=Cm(2.5)
        normal=document.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(11)
        heading=document.add_heading(title or s.DOC_DEFAULT_TITLE,0); heading.alignment=WD_ALIGN_PARAGRAPH.CENTER
        values={s.DOC_SOURCE_FILE:info.filename,s.DOC_GENERATED:generated_at,
                s.DOC_DURATION:format_timestamp(info.duration),s.DOC_MODEL:model or MODEL}
        # Named, the reviewer appears above the transcript as well as beside every comment.
        if (comment_author or "").strip(): values[s.DOC_REVIEWER]=comment_author.strip()
        for key,value in (metadata or {}).items():
            if value: values[key]=value
        for label,value in values.items():
            p=document.add_paragraph(); p.add_run(f"{label}: ").bold=True; p.add_run(str(value))
        if include_notice:
            note=document.add_paragraph(s.DOC_NOTICE); note.runs[0].italic=True
        _legend(document,segments,highlight_labels)
        for item in sorted(segments,key=lambda s:(s.absolute_start,s.absolute_end,s.chunk_index)):
            line=document.add_paragraph(); line.paragraph_format.space_before=Pt(7); line.paragraph_format.space_after=Pt(2)
            speaker_run=line.add_run(apply_speaker_mapping(item,mapping)); speaker_run.bold=True
            if include_timestamps:
                stamp=line.add_run(f"  {format_timestamp(item.absolute_start)}–{format_timestamp(item.absolute_end)}"); stamp.font.color.rgb=RGBColor(100,105,110)
            body=document.add_paragraph(); body.paragraph_format.space_after=Pt(5)
            # Marks on parts of the turn. The paragraph is built one run per stretch of
            # unchanging formatting, so bold and a highlight land on the marked phrase and
            # stop where it stops, and a comment can be anchored to those runs alone.
            anchored: list[tuple[Any,str]] = []
            for piece in an.runs(item.text, getattr(item,"annotations",None) or []) or [an.Run(item.text)]:
                run=body.add_run(piece.text)
                if piece.bold: run.bold=True
                if piece.slot in an.HIGHLIGHTS: _paint(run,piece.slot)
                for text in piece.comments: anchored.append((run,text))
            author=comment_author or s.DOC_COMMENT_AUTHOR
            for run,text in anchored:
                try: document.add_comment([run], text, author=author)
                except Exception: body.add_run(f"  [{text}]").italic=True
            # A note without a range is about the whole turn, so it anchors to all of it.
            note=(getattr(item,"note","") or "").strip()
            if note and body.runs:
                try: document.add_comment(body.runs, note, author=author)
                except Exception: body.add_run(f"  [{note}]").italic=True
        _page_number(section.footer.paragraphs[0]); output=BytesIO(); document.save(output); return output.getvalue()
    except Exception as exc: raise RuntimeError(f"The Word document could not be generated: {exc}") from exc
